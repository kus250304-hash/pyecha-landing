"""
시군구별 자동차 등록대수 통계 (동·구 페이지를 만드는 순서를 정하는 데 쓴다).

우선순위 기준 (docs/roadmap.md 5·6절):
  1순위  시군구별 차령 10년 이상(노후차) 등록대수가 많은 순     → data/vehicle_stats.json        (basis "old10")
         (국토교통부가 시군구별 차령 표를 내지 않아 2026-09 현재 없음)
  2순위  추정 노후 자가용 대수가 많은 순                          → data/vehicle_stats_old10_est.json (basis "old10_est")
         = 시군구별 자가용(관용·영업용 제외) 등록대수 × 그 시도의 자가용 중 차령 10년 이상 비율.
         엑셀에 자가용 칸이 없으면 전체 등록대수 × 시도 전체 노후차 비율(파일의 "private" 가 false).
         `molit` 명령이 국토교통부 월별 엑셀(자동차 등록자료 통계.xlsx)에서 바로 만든다.
  3순위  시군구별 전체 등록대수가 많은 순                         → data/vehicle_stats_total.json  (basis "total")
  4순위  모두 없으면 구 페이지가 있는 구 → 동 페이지가 많은 구 순. 매일 보고에 "통계 없음"을 알린다.

추정 노후 자가용 대수 만들기 (2순위, 매월 한 번):
  python3 scripts/vehicle_stats.py molit                 # 통계누리에서 최신 월 엑셀을 받아 계산·저장
  python3 scripts/vehicle_stats.py molit 받은파일.xlsx   # 이미 받은 파일로 계산
  - 02.통계표_시군구 시트의 "총계" 아래 "자가용"(없으면 "계") 칸 × 14.차종별_상세등록(시도) 시트의
    "합계" 아래 "자가용"(없으면 "계") 칸에서 구한 시도별 비율(모델연도 ≤ 조회연도-10 인 줄의 합 / 총계).
    14번 시트 첫 줄(2006)은 그 해 이전 모델을 모두 포함한 줄이다.
  - openpyxl 이 필요하다(pip install openpyxl).

저장 형식
  {"basis": "old10", "as_of": "2026년 8월", "source": "국토교통부 자동차등록현황보고(차령별)",
   "url": "https://...", "saved_on": "YYYY-MM-DD",
   "rows": [{"sido": "서울특별시", "sigungu": "종로구", "count": 12345}, ...]}

통계 파일 넣기 (공공데이터포털·국토교통부 통계누리에서 받은 CSV 또는 엑셀 그대로):
  python3 scripts/vehicle_stats.py import 받은파일.csv --basis old10 --as-of "2026년 8월" \\
      --source "국토교통부 자동차등록현황보고(차령별)" --url "https://stat.molit.go.kr/..." --count-col "10년이상"
  - 시도·시군구 칸은 이름으로 찾는다(시도/시·도, 시군구/시·군·구). 숫자 칸은 --count-col 로 고르거나,
    --sum-cols "10년,11년,12년" 처럼 여러 칸을 더한다. 같은 시군구 줄이 여러 개면 더한다.
  - 엑셀(.xlsx)은 openpyxl 이 있을 때만 읽는다. 없으면 CSV 로 저장해서 넣는다.
  - 숫자를 지어내거나 추정해서 넣지 않는다. 받은 파일 그대로만 넣는다.

  python3 scripts/vehicle_stats.py show     # 지금 쓰는 기준과 상위 20개 시군구

구·시 페이지 공공 정보 표의 '자동차 등록대수' 줄(2026-10-03, registration_row):
  - 추정치(count)가 아니라 엑셀에 있는 실제 숫자(자가용 registered, 전체 total)만 기준 연월·출처와 함께 쓴다.
  - 빌더(build_gu.py·build_si.py)가 이 파일에서 자동으로 넣는다. 루틴은 등록대수를 웹에서 찾지 않는다.

옛 이름 통계 맞추기(2026-07-01 행정구역 변경 전 자료):
  - 광주광역시·전라남도 → 전남광주통합특별시, 강원도 → 강원특별자치도, 전라북도 → 전북특별자치도
  - 인천 중구 → 제물포구·영종구, 인천 동구 → 제물포구, 인천 서구 → 서해구·검단구.
    나뉜 구는 옛 구 숫자를 새 구 모두에 그대로 쓰고, 합쳐진 구는 더한다. 순서를 정하는 데만 쓰는 근사값이다.
  - 통계에 '수원시'처럼 시 전체만 있고 구(장안구 등)가 없으면 그 시의 구 모두에 시 전체 숫자를 쓴다(근사).
"""
import argparse
import csv
import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OLD10_PATH = ROOT / "data" / "vehicle_stats.json"
EST_PATH = ROOT / "data" / "vehicle_stats_old10_est.json"
TOTAL_PATH = ROOT / "data" / "vehicle_stats_total.json"
MOLIT_META_URL = "https://stat.molit.go.kr/portal/cate/statMetaView.do?hRsId=58"
MOLIT_DOWN_URL = "https://stat.molit.go.kr/portal/common/downLoadFile.do"
BASIS_LABEL = {
    "old10": "차령 10년 이상 노후차",
    "old10_est": "추정 노후 자가용(시군구 자가용 × 시도 노후차 비율)",
    "total": "전체(노후차 통계 없음)",
}

SIDO_ALIAS = {
    "광주광역시": "전남광주통합특별시", "전라남도": "전남광주통합특별시", "전남광주": "전남광주통합특별시",
    "강원도": "강원특별자치도", "전라북도": "전북특별자치도",
    "서울": "서울특별시", "부산": "부산광역시", "대구": "대구광역시", "인천": "인천광역시", "광주": "전남광주통합특별시",
    "대전": "대전광역시", "울산": "울산광역시", "세종": "세종특별자치시", "경기": "경기도", "강원": "강원특별자치도",
    "충북": "충청북도", "충남": "충청남도", "전북": "전북특별자치도", "전남": "전남광주통합특별시",
    "경북": "경상북도", "경남": "경상남도", "제주": "제주특별자치도",
}
# (시도, 옛 시군구) → 새 시군구들
SIGUNGU_ALIAS = {
    ("인천광역시", "중구"): ["제물포구", "영종구"],
    ("인천광역시", "동구"): ["제물포구"],
    ("인천광역시", "서구"): ["서해구", "검단구"],
}


def _key(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def load(prefer: str = "old10") -> tuple[str, dict[tuple[str, str], int], dict]:
    """(기준, {(시도, 시군구 공백없음): 등록대수}, 메타). 파일이 없으면 ("none", {}, {})."""
    for basis, path in (("old10", OLD10_PATH), ("old10_est", EST_PATH), ("total", TOTAL_PATH)):
        if prefer == "total" and basis != "total":
            continue
        if path.exists():
            data = json.loads(path.read_text(encoding="utf-8"))
            if data.get("rows"):
                counts: dict[tuple[str, str], int] = {}
                for r in data["rows"]:
                    k = (r["sido"], _key(r["sigungu"]))
                    counts[k] = counts.get(k, 0) + int(r["count"])
                return data.get("basis", basis), counts, {k: v for k, v in data.items() if k != "rows"}
    return "none", {}, {}


def count_for(counts: dict[tuple[str, str], int], sido: str, sigungu: str) -> int | None:
    """그 시군구의 등록대수. '수원시 장안구' 가 없으면 '수원시' 전체를 쓴다. 없으면 None.
    세종처럼 시군구가 빈칸이면 시도 이름 줄(통계의 '세종특별자치시')을 쓴다.
    CSV 에는 시 하나인데 통계는 구로 나뉘어 있으면 그 시의 구를 모두 더한다(2026-09-30 화성시 구 반영 전에 쓰던 예외, 다른 시가 같은 상태가 되면 다시 쓰인다)."""
    k = _key(sigungu) or _key(sido)
    if k.endswith("시"):
        parts = [n for (s, g), n in counts.items() if s == sido and g.startswith(k) and g != k and g.endswith("구")]
        if parts:
            return sum(parts) + counts.get((sido, k), 0)
    if (sido, k) in counts:
        return counts[(sido, k)]
    city = sigungu.split()[0] if " " in sigungu else None
    if city and (sido, _key(city)) in counts:
        return counts[(sido, _key(city))]
    return None


def _read_table(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            raise SystemExit("엑셀을 읽으려면 openpyxl 이 필요합니다(pip install openpyxl). 아니면 CSV 로 저장해서 넣으세요.")
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
        rows = [[("" if c is None else str(c)).strip() for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp949"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        rows = [[c.strip() for c in row] for row in csv.reader(text.splitlines())]
    # 머리줄: 시도·시군구 칸이 모두 있는 첫 줄
    for i, row in enumerate(rows):
        names = [_key(c) for c in row]
        if any(n in ("시도", "시·도") for n in names) and any(n in ("시군구", "시·군·구") for n in names):
            header = [_key(c).replace("·", "") for c in row]
            return [dict(zip(header, r)) for r in rows[i + 1:] if any(r)]
    raise SystemExit("머리줄에서 '시도'와 '시군구' 칸을 찾지 못했습니다. 칸 이름을 확인해 주세요.")


def _num(s: str) -> int:
    s = re.sub(r"[,\s]", "", s or "")
    return int(float(s)) if re.fullmatch(r"-?\d+(\.\d+)?", s) else 0


def cmd_import(args: argparse.Namespace) -> None:
    table = _read_table(Path(args.file))
    cols = [_key(c) for c in (args.sum_cols.split(",") if args.sum_cols else [args.count_col] if args.count_col else [])]
    if not cols:
        raise SystemExit("--count-col 또는 --sum-cols 로 숫자 칸을 골라 주세요. 칸 이름: " + ", ".join(table[0].keys()))
    missing = [c for c in cols if c not in table[0]]
    if missing:
        raise SystemExit(f"없는 칸: {missing}. 칸 이름: {', '.join(table[0].keys())}")
    rows: dict[tuple[str, str], int] = {}
    for r in table:
        sido, sigungu = r.get("시도", ""), r.get("시군구", "")
        if not sido or not sigungu or sigungu in ("계", "합계", "소계") or sido in ("계", "합계", "전국"):
            continue
        sido = SIDO_ALIAS.get(sido, sido)
        n = sum(_num(r.get(c, "")) for c in cols)
        for new in SIGUNGU_ALIAS.get((sido, sigungu), [sigungu]):
            rows[(sido, new)] = rows.get((sido, new), 0) + n
    if not rows:
        raise SystemExit("읽은 줄이 없습니다")
    out = {
        "basis": args.basis, "as_of": args.as_of, "source": args.source, "url": args.url,
        "saved_on": date.today().isoformat(),
        "rows": [{"sido": s, "sigungu": g, "count": n} for (s, g), n in sorted(rows.items(), key=lambda x: -x[1])],
    }
    path = OLD10_PATH if args.basis == "old10" else TOTAL_PATH
    path.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"저장: {path.relative_to(ROOT)} ({len(rows)}개 시군구, 기준 {args.basis}, {args.as_of})")


def _sheet(wb, prefix: str):
    for ws in wb.worksheets:
        if ws.title.startswith(prefix):
            return [[("" if c is None else c) for c in row] for row in ws.iter_rows(values_only=True)]
    raise SystemExit(f"엑셀에 '{prefix}…' 시트가 없습니다. 파일 형식이 바뀌었는지 확인해 주세요.")


def _group_cols(rows: list, group: str) -> tuple[int, dict[str, int]]:
    """머리줄에서 group(예: '총계') 칸을 찾아 (머리줄 번호, {'자가용': 열, '계': 열, …})."""
    for i, row in enumerate(rows[:10]):
        for j, c in enumerate(row):
            if _key(str(c)) == group:
                sub, end = {}, j + 1
                while end < len(row) and row[end] in ("", None):
                    end += 1
                for jj in range(j, end):
                    sub[_key(str(rows[i + 1][jj]))] = jj
                return i, sub
    raise SystemExit(f"머리줄에서 '{group}' 칸을 찾지 못했습니다.")


def _fetch_latest() -> tuple[Path, str]:
    """통계누리 자동차등록현황보고 화면에서 가장 위(최신) '자동차 등록자료 통계.xlsx' 를 받는다."""
    import tempfile
    import time
    import urllib.parse
    import urllib.request

    def get(url: str) -> bytes:
        for i in range(6):  # 통계누리는 연결이 자주 끊겨 몇 번 다시 시도한다
            try:
                with urllib.request.urlopen(url, timeout=120) as r:
                    return r.read()
            except Exception as e:  # noqa: BLE001
                err = e
                time.sleep(2 + i * 2)
        raise SystemExit(f"통계누리 접속 실패: {err}")

    html = get(MOLIT_META_URL).decode("utf-8", "replace")
    m = re.search(r"downFile\('([^']*자동차 등록자료 통계[^']*\.xlsx)','([^']*)','([^']*)'", html)
    if not m:
        raise SystemExit("통계누리 화면에서 '자동차 등록자료 통계.xlsx' 를 찾지 못했습니다.")
    qs = urllib.parse.urlencode({"oFileName": m.group(1), "rFileName": m.group(2), "midpath": m.group(3)})
    data = get(f"{MOLIT_DOWN_URL}?{qs}")
    if not data.startswith(b"PK"):
        raise SystemExit("받은 파일이 엑셀이 아닙니다.")
    path = Path(tempfile.gettempdir()) / "molit_car_stats.xlsx"
    path.write_bytes(data)
    print(f"받음: {m.group(1)}")
    return path, m.group(1)


def cmd_molit(args: argparse.Namespace) -> None:
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("엑셀을 읽으려면 openpyxl 이 필요합니다(pip install openpyxl).")
    if args.file:
        path, fname = Path(args.file), Path(args.file).name
    else:
        path, fname = _fetch_latest()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # 시군구별 등록대수 (02.통계표_시군구, '총계' 묶음)
    gu = _sheet(wb, "02.")
    ym = str(gu[1][1]).strip()  # 조회년월 '2026.08'
    ymm = re.fullmatch(r"(\d{4})\.(\d{1,2})", ym)
    if not ymm:
        raise SystemExit(f"조회년월을 읽지 못했습니다: {ym!r}")
    year, month = int(ymm.group(1)), int(ymm.group(2))
    hi, cols = _group_cols(gu, "총계")
    private = "자가용" in cols
    ccol = cols["자가용"] if private else cols["계"]
    tcol = cols.get("계")  # 전체 등록대수(구·시 페이지 공공 정보 표에 씀, 2026-10-03)

    # 시도별 노후차 비율 (14.차종별_상세등록(시도), '합계' 묶음, 모델연도별)
    sd = _sheet(wb, "14.")
    shi, scols = _group_cols(sd, "합계")
    scol = scols["자가용"] if private else scols["계"]
    cutoff = year - 10
    old: dict[str, int] = {}
    tot: dict[str, int] = {}
    sido = ""
    for r in sd[shi + 2:]:
        if r[0]:
            sido = str(r[0]).strip()
        y = str(r[1]).strip()
        if not sido or r[scol] in ("", None):
            continue
        if y == "총계":
            tot[sido] = int(r[scol])
        elif re.fullmatch(r"\d{4}", y) and int(y) <= cutoff:
            old[sido] = old.get(sido, 0) + int(r[scol])
    ratio = {s: old.get(s, 0) / tot[s] for s in tot if tot[s]}

    rows: dict[tuple[str, str], list[int]] = {}
    approx: set[tuple[str, str]] = set()  # 옛 구 숫자를 새 구에 그대로 옮긴 줄(근사값, 페이지에는 안 씀)
    sido = ""
    for r in gu[hi + 2:]:
        if r[0]:
            sido = str(r[0]).strip()
        g = str(r[1]).strip()
        if not sido or not g or g in ("계", "합계", "소계") or sido in ("계", "합계", "전국"):
            continue
        if sido not in ratio:
            raise SystemExit(f"'{sido}' 의 노후차 비율을 14번 시트에서 찾지 못했습니다.")
        full = SIDO_ALIAS.get(sido, sido)
        n = int(r[ccol] or 0)
        t = int(r[tcol] or 0) if tcol is not None else 0
        for new in SIGUNGU_ALIAS.get((full, g), [g]):
            cur = rows.setdefault((full, new), [0, 0, 0, 0])
            cur[0] += n
            cur[1] += round(n * ratio[sido])
            cur[2] = round(ratio[sido] * 10000)
            cur[3] += t
            if new != g:
                approx.add((full, new))
    if not rows:
        raise SystemExit("읽은 줄이 없습니다")

    kind = "자가용" if private else "전체"
    out = {
        "basis": "old10_est", "as_of": f"{year}년 {month}월", "private": private,
        "method": f"시군구 {kind} 등록대수 × 시도 {kind} 중 모델연도 {cutoff}년 이하 비율 (반올림)",
        "source": f"국토교통부 자동차등록현황보고({fname}: 02.통계표_시군구, 14.차종별_상세등록(시도))",
        "url": MOLIT_META_URL, "saved_on": date.today().isoformat(),
        "sido_ratio": {SIDO_ALIAS.get(s, s): round(v, 4) for s, v in sorted(ratio.items(), key=lambda x: -x[1])},
        "rows": [{"sido": s, "sigungu": g, "count": v[1], "registered": v[0], "old_ratio": v[2] / 10000,
                  **({"total": v[3]} if tcol is not None else {}), **({"approx": True} if (s, g) in approx else {})}
                 for (s, g), v in sorted(rows.items(), key=lambda x: -x[1][1])],
    }
    EST_PATH.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"저장: {EST_PATH.relative_to(ROOT)} ({len(rows)}개 시군구, {out['as_of']}, {out['method']})")
    if not private:
        print("알림: 엑셀에 자가용 칸이 없어 전체 등록대수 × 시도 노후차 비율로 계산했습니다.")


def _official_rows() -> tuple[dict, dict[tuple[str, str], tuple[int | None, int]]] | None:
    """페이지에 쓸 실제 등록대수: {(시도, 시군구 공백없음): (전체, 자가용)}. 추정치(count·old_ratio)는 쓰지 않는다.
    자가용은 추정 파일의 registered(엑셀 그대로), 전체는 그 파일의 total 또는 같은 달 vehicle_stats_total.json.
    옛 구 숫자를 옮겨 적은 근사 줄(approx)은 뺀다."""
    if not EST_PATH.exists():
        return None
    est = json.loads(EST_PATH.read_text(encoding="utf-8"))
    if not est.get("rows") or not est.get("private"):
        return None
    totals: dict[tuple[str, str], int] = {}
    if TOTAL_PATH.exists():
        tot = json.loads(TOTAL_PATH.read_text(encoding="utf-8"))
        if tot.get("as_of") == est.get("as_of"):
            totals = {(r["sido"], _key(r["sigungu"])): int(r["count"]) for r in tot.get("rows", [])}
    out = {}
    for r in est["rows"]:
        if r.get("approx"):
            continue
        k = (r["sido"], _key(r["sigungu"]))
        out[k] = (r.get("total", totals.get(k)), int(r["registered"]))
    return est, out


def _fmt(total: int | None, private: int) -> str:
    return (f"전체 {total:,}대, 그중 자가용 {private:,}대" if total else f"자가용 {private:,}대")


def registration_row(sido: str, sigungu: str) -> dict | None:
    """구 페이지 공공 정보 표의 '자동차 등록대수' 줄(2026-10-03). 국토교통부 자동차등록현황(월별 엑셀)의 시군구 숫자 그대로.
    - 그 시군구 줄이 있으면 그 숫자.
    - '천안시 서북구'처럼 구 숫자가 없는 일반구는 그 시 전체(통계의 그 시 줄을 모두 더한 값)를 '○○시 전체 기준'으로 밝혀 쓴다. 구 숫자를 추정하지 않는다.
    - 통계에 없으면 None(줄을 넣지 않음)."""
    got = _official_rows()
    if not got:
        return None
    est, rows = got
    as_of = est["as_of"]
    k = _key(sigungu) or _key(sido)
    if (sido, k) in rows:
        total, private = rows[(sido, k)]
        value = f"{_fmt(total, private)} ({as_of} 기준)"
    else:
        city = sigungu.split()[0] if " " in sigungu else ""
        parts = [v for (s, g), v in rows.items() if city and s == sido and (g == _key(city) or (g.startswith(_key(city)) and g.endswith("구")))]
        if not parts:
            return None
        total = sum(t for t, _ in parts) if all(t for t, _ in parts) else None
        value = f"{city} 전체 기준({as_of}): {_fmt(total, sum(p for _, p in parts))}. 통계에 구별 숫자가 없어 시 전체 숫자를 적습니다"
    return {"label": "자동차 등록대수", "value": value, "auto": True,
            "source": f"국토교통부 자동차등록현황보고({as_of} 자동차 등록자료 통계, 02.통계표_시군구)", "url": est.get("url") or MOLIT_META_URL}


def registration_row_sido(sido: str) -> dict | None:
    """시·도 페이지의 '자동차 등록대수' 줄: 그 시·도 시군구 줄을 모두 더한 값. 근사 줄이 섞이면 넣지 않는다."""
    got = _official_rows()
    if not got:
        return None
    est, rows = got
    if any(r.get("approx") and r["sido"] == sido for r in est["rows"]):
        return None
    parts = [v for (s, _), v in rows.items() if s == sido]
    if not parts:
        return None
    total = sum(t for t, _ in parts) if all(t for t, _ in parts) else None
    return {"label": "자동차 등록대수", "value": f"{_fmt(total, sum(p for _, p in parts))} ({est['as_of']} 기준, 시·군·구 통계 합계)", "auto": True,
            "source": f"국토교통부 자동차등록현황보고({est['as_of']} 자동차 등록자료 통계, 02.통계표_시군구)", "url": est.get("url") or MOLIT_META_URL}


def with_registration(info: list[dict], row: dict | None) -> list[dict]:
    """공공 정보 줄 목록에 등록대수 줄을 붙인다. 데이터에 이미 등록대수 줄이 있으면 그대로 둔다."""
    if not row or any("등록대수" in r.get("label", "") for r in info):
        return list(info)
    return list(info) + [row]


def cmd_show(_: argparse.Namespace) -> None:
    basis, counts, meta = load()
    if basis == "none":
        print("통계 없음: data/vehicle_stats.json(노후차)·vehicle_stats_total.json(전체) 둘 다 없음 → 동 페이지 많은 구 순으로 뽑음")
        return
    print(f"기준: {BASIS_LABEL.get(basis, basis)} 등록대수 · {meta.get('as_of')} · {meta.get('source')}")
    if meta.get("method"):
        print(f"계산: {meta['method']}")
    # 동을 뽑을 때와 같은 단위(법정동 목록의 시군구)로 센다. 목록에 구가 없는 시는 통계의 구를 합친다.
    with (ROOT / "data" / "legal_dong_list.csv").open(encoding="utf-8-sig") as f:
        units = sorted({(r["시도"], r["시군구"]) for r in csv.DictReader(f)})
    ranked = sorted(((n, s, g) for s, g in units if (n := count_for(counts, s, g)) is not None), reverse=True)
    for i, (n, s, g) in enumerate(ranked[:20], 1):
        print(f"  {i:2d}. {s} {g or s}: {n:,}")
    missing = [f"{s} {g}".strip() for s, g in units if count_for(counts, s, g) is None]
    if missing:
        print(f"  통계에 없는 시군구 {len(missing)}곳(맨 뒤로 감): {', '.join(missing[:10])}")


def main() -> None:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("import")
    p.add_argument("file")
    p.add_argument("--basis", choices=["old10", "total"], required=True)
    p.add_argument("--as-of", required=True, help='통계 기준 연월, 예: "2026년 8월"')
    p.add_argument("--source", required=True)
    p.add_argument("--url", required=True)
    p.add_argument("--count-col")
    p.add_argument("--sum-cols")
    p.set_defaults(func=cmd_import)
    p = sub.add_parser("molit", help="국토교통부 월별 엑셀로 추정 노후 자가용 대수 계산")
    p.add_argument("file", nargs="?", help="받은 엑셀 파일 (없으면 통계누리에서 최신 파일을 받음)")
    p.set_defaults(func=cmd_molit)
    sub.add_parser("show").set_defaults(func=cmd_show)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
